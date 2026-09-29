"""
디플로맷 연수회 PPT 양식 고도화 빌더 (build_reference_styled_presentation.py)
반영 사항:
1. 대제목(24pt) 및 소제목(16pt) 글자 크기 엄격 유지
2. 한 줄을 넘어가는 문장은 \r\n을 적용하여 어절 단위로 다음 줄에 깔끔하게 시작되도록 정리
3. 11장 압축 구성 및 목차-본문 대제목 1:1 매핑
4. 목차에 네모(□) 없는 순수 텍스트 리스트
5. ./reference 스타일 (파스텔 3카드, 헤더 밴드 카드, 6단계 파이프라인 플로우) 완벽 적용
"""

import sys
import os
from pathlib import Path
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

if sys.platform == "win32":
    sys.stdout.reconfigure(encoding='utf-8')

# Diplomat 연수회 표준 색상 팔레트 (from diplomat-연수회양식 skill)
COLOR_TEAL_PRIMARY   = RGBColor(0x2D, 0xA2, 0xBF)  # Accent 1: Diplomat Teal (#2DA2BF)
COLOR_RED_ACCENT     = RGBColor(0xDA, 0x1F, 0x28)  # Accent 2: Crimson Red (#DA1F28)
COLOR_ORANGE_ACCENT  = RGBColor(0xEB, 0x64, 0x1B)  # Accent 3: Orange (#EB641B)
COLOR_NAVY_ACCENT    = RGBColor(0x39, 0x63, 0x9D)  # Accent 4: Navy Blue (#39639D)
COLOR_INDIGO_NAVY    = RGBColor(0x47, 0x4B, 0x78)  # Accent 5: Indigo Navy (#474B78)
COLOR_WINE_MAROON    = RGBColor(0x7D, 0x3C, 0x4A)  # Accent 6: Wine Maroon (#7D3C4A)
COLOR_TEXT_DARK      = RGBColor(0x46, 0x46, 0x46)  # Dark 2: Charcoal Gray (#464646)
COLOR_ICE_BLUE       = RGBColor(0xDE, 0xF5, 0xFA)  # Light 2: Soft Ice Blue (#DEF5FA)
COLOR_WHITE          = RGBColor(0xFF, 0xFF, 0xFF)  # Background: Pure White (#FFFFFF)

# Diplomat 표준 카드 배경 및 테두리 (스킬 테마 연계)
# 1. Teal 카드 (주요 개념, 도구, 긍정)
COLOR_TEAL_CARD_BG     = RGBColor(0xDE, 0xF5, 0xFA)  # Soft Ice Blue (#DEF5FA)
COLOR_TEAL_CARD_BORDER = RGBColor(0x2D, 0xA2, 0xBF)  # Diplomat Teal (#2DA2BF)
COLOR_TEAL_CARD_TITLE  = RGBColor(0x2D, 0xA2, 0xBF)

# 2. Navy 카드 (시스템, 에이전트, 신뢰)
COLOR_NAVY_CARD_BG     = RGBColor(0xEB, 0xF1, 0xF8)  # Soft Navy tint (#EBF1F8)
COLOR_NAVY_CARD_BORDER = RGBColor(0x39, 0x63, 0x9D)  # Navy Blue (#39639D)
COLOR_NAVY_CARD_TITLE  = RGBColor(0x39, 0x63, 0x9D)

# 3. Orange 카드 (실천, 플러그인, 행동 유발)
COLOR_ORANGE_CARD_BG     = RGBColor(0xFD, 0xF0, 0xE8)  # Soft Orange tint (#FDF0E8)
COLOR_ORANGE_CARD_BORDER = RGBColor(0xEB, 0x64, 0x1B)  # Orange (#EB641B)
COLOR_ORANGE_CARD_TITLE  = RGBColor(0xEB, 0x64, 0x1B)

# 4. Red 카드 (경고, 문제점, Before)
COLOR_RED_CARD_BG     = RGBColor(0xFD, 0xF2, 0xF2)  # Soft Red tint (#FDF2F2)
COLOR_RED_CARD_BORDER = RGBColor(0xDA, 0x1F, 0x28)  # Crimson Red (#DA1F28)
COLOR_RED_CARD_TITLE  = RGBColor(0xDA, 0x1F, 0x28)

COLOR_GRAY_LIGHT  = RGBColor(0xFA, 0xFC, 0xFD)  # 매우 연한 서브 배경
COLOR_GRAY_BORDER = RGBColor(0xD2, 0xEB, 0xF1)  # Soft Teal-Gray border

FONT_FAMILY = "맑은 고딕"
TEMPLATE_PATH = Path(r"C:\Users\user\.gemini\config\skills\diplomat-연수회양식\resources\template.pptx")


class DiplomatReferenceStyleBuilder:
    def __init__(self, template_path=None):
        t_path = Path(template_path) if template_path else TEMPLATE_PATH
        if not t_path.exists():
            raise FileNotFoundError(f"Template not found: {t_path}")
        self.prs = Presentation(str(t_path))
        self._clear_slides()
        self.page_counter = 1

    def _clear_slides(self):
        for i in range(len(self.prs.slides) - 1, -1, -1):
            rId = self.prs.slides._sldIdLst[i].rId
            self.prs.part.drop_rel(rId)
            del self.prs.slides._sldIdLst[i]

    def _setup_header(self, slide, title, key_message="", page_num=None):
        """대제목(24pt)과 소제목(16pt) 세팅 - 글자 크기 유지 및 \r\n 줄바꿈 반영"""
        cur_num = page_num if page_num is not None else self.page_counter
        for ph in slide.placeholders:
            idx = ph.placeholder_format.idx
            if idx == 10:
                # 대제목: 24pt Bold 유지
                lines = title.replace("\r\n", "\n").split("\n")
                tf = ph.text_frame
                tf.word_wrap = True
                for i, line in enumerate(lines):
                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    p.text = line
                    p.font.name = FONT_FAMILY
                    p.font.size = Pt(24)
                    p.font.bold = True
                    p.font.color.rgb = COLOR_TEXT_DARK
            elif idx in (13, 14):
                # 소제목: 16pt Bold Teal 유지 (1줄 단일행으로 깔끔하게 표기)
                clean_msg = key_message.replace("\r\n", " ").replace("\n", " ").strip()
                tf = ph.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.text = clean_msg
                p.font.name = FONT_FAMILY
                p.font.size = Pt(16)
                p.font.bold = True
                p.font.color.rgb = COLOR_TEAL_PRIMARY
            elif idx == 12:
                ph.text = str(cur_num)
        self.page_counter += 1

    def _add_diplomat_callout_box(self, slide, text):
        """디플로맷 표준 Soft Ice Blue (#DEF5FA) 배경 + Teal 테두리 하단 요약 박스"""
        if not text:
            return
        box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.65), Inches(5.45), Inches(8.70), Inches(0.68)
        )
        box.fill.solid()
        box.fill.fore_color.rgb = COLOR_ICE_BLUE
        box.line.color.rgb = COLOR_TEAL_PRIMARY
        box.line.width = Pt(1.2)

        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.12)
        tf.margin_bottom = Inches(0.10)
        tf.margin_left = Inches(0.20)
        tf.margin_right = Inches(0.20)

        lines = text.replace("\r\n", "\n").split("\n")
        for j, line in enumerate(lines):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.text = line
            p.alignment = PP_ALIGN.CENTER
            p.font.name = FONT_FAMILY
            p.font.size = Pt(13 if len(lines) > 1 else 14)
            p.font.bold = True
            p.font.color.rgb = COLOR_TEXT_DARK

    # 1. 표지 (Layout 0)
    def add_title_slide(self, title, date_str, team, presenter, header_tag="수요연수회"):
        layout = self.prs.slide_layouts[0]
        slide = self.prs.slides.add_slide(layout)
        for ph in slide.placeholders:
            idx = ph.placeholder_format.idx
            if idx == 10:
                ph.text = title
                for p in ph.text_frame.paragraphs:
                    p.font.name = FONT_FAMILY
                    p.font.size = Pt(28)
                    p.font.bold = True
                    p.font.color.rgb = COLOR_TEXT_DARK
            elif idx == 11:
                ph.text = date_str
                for p in ph.text_frame.paragraphs:
                    p.font.name = FONT_FAMILY
                    p.font.size = Pt(14)
                    p.font.color.rgb = COLOR_TEXT_DARK
            elif idx == 12:
                ph.text = team
                for p in ph.text_frame.paragraphs:
                    p.font.name = FONT_FAMILY
                    p.font.size = Pt(14)
                    p.font.color.rgb = COLOR_TEXT_DARK
            elif idx == 13:
                ph.text = presenter
                for p in ph.text_frame.paragraphs:
                    p.font.name = FONT_FAMILY
                    p.font.size = Pt(14)
                    p.font.color.rgb = COLOR_TEXT_DARK

        if header_tag:
            tx_box = slide.shapes.add_textbox(Inches(0.42), Inches(0.51), Inches(3.5), Inches(0.6))
            tf = tx_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = header_tag
            p.font.name = FONT_FAMILY
            p.font.size = Pt(18)
            p.font.bold = True
            p.font.color.rgb = COLOR_TEAL_PRIMARY
        self.page_counter += 1
        return slide

    # 2. 목차 (Layout 1) - 네모(□) 및 숫자(1. 2.) 없이 깔끔한 텍스트 리스트
    def add_agenda_slide(self, items, title="목차"):
        layout = self.prs.slide_layouts[1]
        slide = self.prs.slides.add_slide(layout)
        for ph in slide.placeholders:
            idx = ph.placeholder_format.idx
            if idx == 10:
                ph.text = title
                for p in ph.text_frame.paragraphs:
                    p.font.name = FONT_FAMILY
                    p.font.size = Pt(28)
                    p.font.bold = True
                    p.font.color.rgb = COLOR_TEXT_DARK
            elif idx == 11:
                tf = ph.text_frame
                tf.word_wrap = True
                for i, item in enumerate(items):
                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    p.text = item
                    p.font.name = FONT_FAMILY
                    p.font.size = Pt(18)
                    p.font.bold = True
                    p.font.color.rgb = COLOR_TEXT_DARK
                    p.space_after = Pt(22)
        self.page_counter += 1
        return slide

    # [디자인 패턴 1] Reference 2: 3단 라운드 카드 + 하단 디플로맷 요약
    def add_pastel_3cards_slide(self, title, key_message, cards_data, bottom_callout):
        layout = self.prs.slide_layouts[2]
        slide = self.prs.slides.add_slide(layout)
        self._setup_header(slide, title, key_message)

        lefts = [Inches(0.65), Inches(3.68), Inches(6.71)]
        card_w = Inches(2.64)
        card_t = Inches(1.85)
        card_h = Inches(3.40)

        for i, c in enumerate(cards_data):
            shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, lefts[i], card_t, card_w, card_h)
            shape.fill.solid()
            shape.fill.fore_color.rgb = c['bg']
            shape.line.color.rgb = c['border']
            shape.line.width = Pt(1.5)

            tf = shape.text_frame
            tf.word_wrap = True
            tf.margin_top = Inches(0.40)
            tf.margin_bottom = Inches(0.2)
            tf.margin_left = Inches(0.15)
            tf.margin_right = Inches(0.15)

            # 제목
            lines_title = c['title'].replace("\r\n", "\n").split("\n")
            for j, t_line in enumerate(lines_title):
                p0 = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                p0.text = t_line
                p0.alignment = PP_ALIGN.CENTER
                p0.font.name = FONT_FAMILY
                p0.font.size = Pt(20 if len(lines_title) > 1 else 22)
                p0.font.bold = True
                p0.font.color.rgb = c['title_color']
            p0.space_after = Pt(18)

            # 본문 설명 (핵심 요약)
            lines_body = c['body'].replace("\r\n", "\n").split("\n")
            for j, b_line in enumerate(lines_body):
                p1 = tf.add_paragraph()
                p1.text = b_line
                p1.alignment = PP_ALIGN.CENTER
                p1.font.name = FONT_FAMILY
                p1.font.size = Pt(15)
                p1.font.bold = True
                p1.font.color.rgb = COLOR_TEXT_DARK
            p1.space_after = Pt(18)

            # 서브 설명 (예시)
            lines_sub = c['sub'].replace("\r\n", "\n").split("\n")
            for j, s_line in enumerate(lines_sub):
                p2 = tf.add_paragraph()
                p2.text = s_line
                p2.alignment = PP_ALIGN.CENTER
                p2.font.name = FONT_FAMILY
                p2.font.size = Pt(12)
                p2.font.color.rgb = COLOR_TEXT_DARK

        # 하단 요약 문구 (디플로맷 표준 Soft Ice Blue 박스)
        self._add_diplomat_callout_box(slide, bottom_callout)
        return slide

    # [디자인 패턴 2] Reference 1: 상단 컬러 띠 카드 (Header-Band Cards)
    def add_header_band_cards_slide(self, title, key_message, cards_data, bottom_callout):
        layout = self.prs.slide_layouts[2]
        slide = self.prs.slides.add_slide(layout)
        self._setup_header(slide, title, key_message)

        num_cards = len(cards_data)
        spacing = Inches(0.25)
        total_w = Inches(8.70)
        card_w = (total_w - spacing * (num_cards - 1)) / num_cards
        start_left = Inches(0.65)
        top_y = Inches(1.85)

        for i, c in enumerate(cards_data):
            c_left = start_left + i * (card_w + spacing)

            # 1) 헤더 띠
            header_h = Inches(0.75)
            h_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_left, top_y, card_w, header_h)
            h_shape.fill.solid()
            h_shape.fill.fore_color.rgb = c['header_color']
            h_shape.line.fill.background()
            tf_h = h_shape.text_frame
            tf_h.word_wrap = True
            lines_ht = c['title'].replace("\r\n", "\n").split("\n")
            for j, h_line in enumerate(lines_ht):
                p_h = tf_h.paragraphs[0] if j == 0 else tf_h.add_paragraph()
                p_h.text = h_line
                p_h.alignment = PP_ALIGN.CENTER
                p_h.font.name = FONT_FAMILY
                p_h.font.size = Pt(14 if len(lines_ht) > 1 else 15)
                p_h.font.bold = True
                p_h.font.color.rgb = COLOR_WHITE

            # 2) 바디 박스
            body_top = top_y + Inches(0.65)
            body_h = Inches(2.65)
            b_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_left, body_top, card_w, body_h)
            b_shape.fill.solid()
            b_shape.fill.fore_color.rgb = COLOR_WHITE
            b_shape.line.color.rgb = c.get('border_color', c['header_color'])
            b_shape.line.width = Pt(1.2)

            tf_b = b_shape.text_frame
            tf_b.word_wrap = True
            tf_b.margin_top = Inches(0.35)
            tf_b.margin_left = Inches(0.15)
            tf_b.margin_right = Inches(0.15)
            lines_b = c['body'].replace("\r\n", "\n").split("\n")
            for j, b_line in enumerate(lines_b):
                p_b = tf_b.paragraphs[0] if j == 0 else tf_b.add_paragraph()
                p_b.text = b_line
                p_b.alignment = PP_ALIGN.CENTER
                p_b.font.name = FONT_FAMILY
                p_b.font.size = Pt(13)
                p_b.font.bold = True
                p_b.font.color.rgb = COLOR_TEXT_DARK

        # 하단 요약 문구 (디플로맷 표준 Soft Ice Blue 박스)
        self._add_diplomat_callout_box(slide, bottom_callout)
        return slide

    # [디자인 패턴 3] Reference 3: 파이프라인 플로우 (Pill Badges & Flow)
    def add_pipeline_flow_slide(self, title, key_message, steps_data, bottom_callout):
        layout = self.prs.slide_layouts[2]
        slide = self.prs.slides.add_slide(layout)
        self._setup_header(slide, title, key_message)

        num_steps = len(steps_data)
        start_left = Inches(0.65)
        arrow_w = Inches(0.20)
        total_content_w = Inches(8.70)
        total_arrow_w = arrow_w * (num_steps - 1)
        badge_w = (total_content_w - total_arrow_w) / num_steps

        badge_top = Inches(2.05)
        badge_h = Inches(0.65)
        card_top = Inches(2.90)
        card_h = Inches(2.35)

        for i, s in enumerate(steps_data):
            cur_left = start_left + i * (badge_w + arrow_w)

            # 1) 알약형 뱃지
            badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cur_left, badge_top, badge_w, badge_h)
            badge.fill.solid()
            badge.fill.fore_color.rgb = s['badge_bg']
            badge.line.color.rgb = s.get('border_color', s['badge_bg'])
            badge.line.width = Pt(1)

            tf = badge.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = s['step']
            p.alignment = PP_ALIGN.CENTER
            p.font.name = FONT_FAMILY
            p.font.size = Pt(12)
            p.font.bold = True
            p.font.color.rgb = s['text_color']

            # 화살표
            if i < num_steps - 1:
                arrow_box = slide.shapes.add_textbox(cur_left + badge_w, badge_top, arrow_w, badge_h)
                tf_a = arrow_box.text_frame
                p_a = tf_a.paragraphs[0]
                p_a.text = "→"
                p_a.alignment = PP_ALIGN.CENTER
                p_a.font.name = FONT_FAMILY
                p_a.font.size = Pt(13)
                p_a.font.bold = True
                p_a.font.color.rgb = COLOR_TEAL_PRIMARY

            # 2) 하단 세부 요약 카드
            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cur_left, card_top, badge_w, card_h)
            card.fill.solid()
            card.fill.fore_color.rgb = COLOR_WHITE
            card.line.color.rgb = COLOR_ICE_BLUE
            card.line.width = Pt(1.5)

            tf_c = card.text_frame
            tf_c.word_wrap = True
            tf_c.margin_top = Inches(0.25)
            tf_c.margin_left = Inches(0.08)
            tf_c.margin_right = Inches(0.08)
            lines_desc = s['desc'].replace("\r\n", "\n").split("\n")
            for j, d_line in enumerate(lines_desc):
                p_c = tf_c.paragraphs[0] if j == 0 else tf_c.add_paragraph()
                p_c.text = d_line
                p_c.alignment = PP_ALIGN.CENTER
                p_c.font.name = FONT_FAMILY
                p_c.font.size = Pt(11)
                p_c.font.bold = True
                p_c.font.color.rgb = COLOR_TEXT_DARK

        # 하단 요약 문구 (디플로맷 표준 Soft Ice Blue 박스)
        self._add_diplomat_callout_box(slide, bottom_callout)
        return slide

    # [디자인 패턴 4] Before vs After 대비 카드 (요약형)
    def add_comparison_slide(self, title, key_message, left_card, right_card, bottom_callout):
        layout = self.prs.slide_layouts[2]
        slide = self.prs.slides.add_slide(layout)
        self._setup_header(slide, title, key_message)

        card_w = Inches(4.20)
        card_h = Inches(3.40)
        card_top = Inches(1.85)
        left_pos = [Inches(0.65), Inches(5.15)]

        cards = [left_card, right_card]
        for i, c in enumerate(cards):
            shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos[i], card_top, card_w, card_h)
            shape.fill.solid()
            shape.fill.fore_color.rgb = c['bg']
            shape.line.color.rgb = c['border']
            shape.line.width = Pt(1.5)

            tf = shape.text_frame
            tf.word_wrap = True
            tf.margin_top = Inches(0.30)
            tf.margin_left = Inches(0.25)
            tf.margin_right = Inches(0.25)

            p0 = tf.paragraphs[0]
            p0.text = c['title']
            p0.font.name = FONT_FAMILY
            p0.font.size = Pt(17)
            p0.font.bold = True
            p0.font.color.rgb = c['title_color']
            p0.space_after = Pt(16)

            for line in c['lines']:
                p = tf.add_paragraph()
                p.text = line
                p.font.name = FONT_FAMILY
                p.font.size = Pt(13)
                p.font.bold = (line.startswith("•") or line.startswith("지시:") or line.startswith("결과:"))
                p.font.color.rgb = COLOR_TEXT_DARK
                p.space_after = Pt(10)

        # 하단 요약 문구 (디플로맷 표준 Soft Ice Blue 박스)
        self._add_diplomat_callout_box(slide, bottom_callout)
        return slide

    # 4. 마무리 슬라이드 (Layout 4)
    def add_closing_slide(self, message="경청해 주셔서 감사합니다.\n\n( 질의응답 및 자유 토론 )"):
        layout = self.prs.slide_layouts[4]
        slide = self.prs.slides.add_slide(layout)
        for ph in slide.placeholders:
            if ph.placeholder_format.idx == 10:
                tf = ph.text_frame
                tf.word_wrap = True
                lines = message.replace("\r\n", "\n").split("\n")
                for j, line in enumerate(lines):
                    p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                    p.text = line
                    p.font.name = FONT_FAMILY
                    p.font.size = Pt(32 if j == 0 else 18)
                    p.font.bold = True
                    p.font.color.rgb = COLOR_TEAL_PRIMARY
                    p.alignment = PP_ALIGN.CENTER
        return slide

    def save(self, output_path):
        out_p = Path(output_path).resolve()
        out_p.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(out_p))
        print(f"[Diplomat Reference Style PPTX] Saved to: {out_p}")
        return out_p


def build_full_presentation():
    builder = DiplomatReferenceStyleBuilder()

    # Slide 1: 표지
    builder.add_title_slide(
        title="AI 활용, 한 단계 더",
        date_str="2026년 9월 30일",
        team="기술연구소",
        presenter="성진영 프로",
        header_tag="제OOO회 수요연수회"
    )

    # 5대 목차 아젠다 (각 핵심 본문 페이지의 대제목과 1:1 완벽 일치)
    agenda_titles = [
        "AI 활용, 어디까지 왔을까?",
        "어떻게 하면 AI를 진짜 잘 쓸 수 있을까?",
        "라이브 시연: 회의 메모의 기적",
        "내일부터 당장 써먹는 1가지 작은 실천",
        "지속 가능한 AI 협업 워크플로우"
    ]

    # Slide 2: 목차 (네모 □ 및 숫자 1. 2. 없는 순수 텍스트 리스트)
    builder.add_agenda_slide(agenda_titles)

    # =========================================================
    # Section 1. AI 활용, 어디까지 왔을까?
    # =========================================================
    # Slide 3: 영상 기술 및 산업 현장 트렌드 (디플로맷 헤더 밴드 3카드: Navy, Teal, Orange)
    builder.add_header_band_cards_slide(
        title=agenda_titles[0],
        key_message="텍스트 한 줄로 시공간을 넘나드는 고품질 영상을 구현하는 시대",
        cards_data=[
            {
                'header_color': COLOR_NAVY_ACCENT,
                'border_color': COLOR_NAVY_ACCENT,
                'title': '영상 사례 1:\r\n타임머신 여행',
                'body': '텍스트 프롬프트 기반\n역사적 시공간 초고속 복원\n\n사실적 영상 자동 생성\n(유튜브 쇼츠 실습 영상)'
            },
            {
                'header_color': COLOR_TEAL_PRIMARY,
                'border_color': COLOR_TEAL_PRIMARY,
                'title': '영상 사례 2:\r\n산업 현장 AI',
                'body': '도면·사진 3D 변환\n회의록 화자 분리 및 요약\n\n단순 작업 즉각 자동화\n(업무 생산성 혁신)'
            },
            {
                'header_color': COLOR_ORANGE_ACCENT,
                'border_color': COLOR_ORANGE_ACCENT,
                'title': '임직원을 위한\r\n시사점',
                'body': '연구용 기술 탈피\n\n실무에 즉시 투입 가능한\n일상 업무 가속 도구'
            }
        ],
        bottom_callout="핵심: AI는 상상을 시각화하고 단순 업무를 초고속화하는 실전 도구입니다."
    )

    # Slide 4: 우리가 AI를 쓰며 느끼는 흔한 좌절 (디플로맷 3카드: Teal, Navy, Orange)
    builder.add_pastel_3cards_slide(
        title="우리가 AI를 쓰며 느끼는 흔한 좌절",
        key_message="기대와 다른 두루뭉술한 답변과 매번 반복되는 설명에 지치는 현실",
        cards_data=[
            {
                'bg': COLOR_TEAL_CARD_BG, 'border': COLOR_TEAL_CARD_BORDER, 'title_color': COLOR_TEAL_CARD_TITLE,
                'title': '두루뭉술한 답변',
                'body': '교과서적 일반론 나열\n\n사내 맥락·규정 미반영',
                'sub': '“알아서 써줘”의 한계'
            },
            {
                'bg': COLOR_NAVY_CARD_BG, 'border': COLOR_NAVY_CARD_BORDER, 'title_color': COLOR_NAVY_CARD_TITLE,
                'title': '반복 설명의 피로',
                'body': '새 대화창마다 복사·붙여넣기\n\n서식·규칙 재설명 피로',
                'sub': '세션 종료 시 기억 망각'
            },
            {
                'bg': COLOR_ORANGE_CARD_BG, 'border': COLOR_ORANGE_CARD_BORDER, 'title_color': COLOR_ORANGE_CARD_TITLE,
                'title': '실무 적용 포기',
                'body': '“그냥 내가 쓰고 말지”\n\n결국 수작업으로 회귀',
                'sub': '도구 활용의 단절'
            }
        ],
        bottom_callout="원인은 AI의 한계가 아닌, '일 시키는 방식(디렉팅)'의 부재였습니다."
    )

    # =========================================================
    # Section 2. 어떻게 하면 AI를 진짜 잘 쓸 수 있을까?
    # =========================================================
    # Slide 5: 3C 원칙 (디플로맷 헤더 밴드 3카드: Navy, Teal, Orange)
    builder.add_header_band_cards_slide(
        title=agenda_titles[1],
        key_message="AI는 자판기가 아니라, 사내 맥락을 모르는 '유능한 신입 사원'입니다.",
        cards_data=[
            {
                'header_color': COLOR_NAVY_ACCENT,
                'border_color': COLOR_NAVY_ACCENT,
                'title': '배경\r\n(Context)',
                'body': '작성 목적 & 보고 대상\n\n사내 상황과 타깃을\n명확히 사전 공유'
            },
            {
                'header_color': COLOR_TEAL_PRIMARY,
                'border_color': COLOR_TEAL_PRIMARY,
                'title': '제약\r\n(Constraints)',
                'body': '필수 포함 & 금지 기준\n\n분량 제한·핵심 수치·\n주의사항 사전 지정'
            },
            {
                'header_color': COLOR_ORANGE_ACCENT,
                'border_color': COLOR_ORANGE_ACCENT,
                'title': '완료 기준\r\n(Format)',
                'body': '최종 출력 양식 규격화\n\n3줄 요약·표 형식 지정\n(AI 자체 검증 유도)'
            }
        ],
        bottom_callout="핵심은 배경·제약·완료조건 3가지를 명확히 분리하여 지시하는 것입니다."
    )

    # Slide 6: 프롬프트 비포/애프터 비교 (Red vs Teal)
    builder.add_comparison_slide(
        title="프롬프트 작성 전과 후 비교 예시",
        key_message="질문의 해상도를 높이면 AI의 답변 수준이 즉시 전문가 수준으로 변합니다.",
        left_card={
            'bg': COLOR_RED_CARD_BG,
            'border': COLOR_RED_ACCENT,
            'title_color': COLOR_RED_ACCENT,
            'title': '❌ 기존 방식 (자판기형 단순 질문)',
            'lines': [
                '• 지시: "신제품 금고 홍보 문구 써줘"',
                '• 문제: 타깃·채널·제약조건 부재',
                '• 결과: 뻔한 교과서적 홍보글 출력',
                '• 평가: 실무 활용 불가 (재작업 발생)'
            ]
        },
        right_card={
            'bg': COLOR_TEAL_CARD_BG,
            'border': COLOR_TEAL_PRIMARY,
            'title_color': COLOR_TEAL_PRIMARY,
            'title': '⭕ 개선 방식 (신입 사원 디렉팅 - 3C 적용)',
            'lines': [
                '• 지시: "30대 1인가구 타깃 방화금고 카피 3개, 안도감 강조, 50자 표로"',
                '• 특징: 타깃·채널·완료기준 명확히 제시',
                '• 결과: 감성을 자극하는 실무용 카피 도출',
                '• 평가: 즉시 결재 가능한 완성도 확보'
            ]
        },
        bottom_callout="질문의 해상도를 높이면, AI는 뻔한 글 대신 실무 완성품을 만듭니다."
    )

    # Slide 7: 스킬 / 에이전트 / 플러그인 (디플로맷 3카드: Teal, Navy, Orange)
    builder.add_pastel_3cards_slide(
        title="한 단계 더: 단발성 질문을 넘어 시스템으로",
        key_message="도구(Tool), 스킬(Skill), 에이전트(Agent)의 결합으로 완성하는 자동화",
        cards_data=[
            {
                'bg': COLOR_TEAL_CARD_BG, 'border': COLOR_TEAL_CARD_BORDER, 'title_color': COLOR_TEAL_CARD_TITLE,
                'title': '스킬 (Skill)',
                'body': '반복 업무 표준 매뉴얼\n\n규칙 영구 기억 (SOP)',
                'sub': '예: 디프로매트 보고서 양식'
            },
            {
                'bg': COLOR_NAVY_CARD_BG, 'border': COLOR_NAVY_CARD_BORDER, 'title_color': COLOR_NAVY_CARD_TITLE,
                'title': '에이전트 (Agent)',
                'body': '목표를 완수하는 전담 비서\n\n스스로 작업 수행',
                'sub': '예: 회의록 요약, 품질 분석 담당'
            },
            {
                'bg': COLOR_ORANGE_CARD_BG, 'border': COLOR_ORANGE_CARD_BORDER, 'title_color': COLOR_ORANGE_CARD_TITLE,
                'title': '플러그인 (Plugin)',
                'body': '외부 도구·기능 연결 연장\n\nAI의 실행 손발',
                'sub': '예: 파일 읽기·쓰기, 검색 도구'
            }
        ],
        bottom_callout="비유하면: 스킬은 업무 매뉴얼, 에이전트는 전담 비서, 플러그인은 손발 도구입니다."
    )

    # =========================================================
    # Section 3. 라이브 시연: 회의 메모의 기적
    # =========================================================
    # Slide 8: 실전 시연 3단계 공정 (디플로맷 헤더 밴드 3카드: Navy, Teal, Orange)
    builder.add_header_band_cards_slide(
        title=agenda_titles[2],
        key_message="부서 간 회의 메모를 투입하여 1분 만에 경영진 보고 양식으로 자동 구조화",
        cards_data=[
            {
                'header_color': COLOR_NAVY_ACCENT,
                'border_color': COLOR_NAVY_ACCENT,
                'title': '1단계:\r\n날것의 메모 투입',
                'body': '두서없는 10줄 회의 메모\n\n부서별 발언·일정 메모 입력\n(현장 날것의 텍스트 원본)'
            },
            {
                'header_color': COLOR_TEAL_PRIMARY,
                'border_color': COLOR_TEAL_PRIMARY,
                'title': '2단계:\r\n사내 표준 스킬 호출',
                'body': '디프로매트 표준 양식 적용\n\n3줄 요약 + 쟁점 표 +\nAction Item 자동 매핑'
            },
            {
                'header_color': COLOR_ORANGE_ACCENT,
                'border_color': COLOR_ORANGE_ACCENT,
                'title': '3단계:\r\n1분 내 보고서 완성',
                'body': '결재용 1페이지 보고서 출력\n\n사람은 1분 팩트체크 후\n즉시 사내 결재 진행'
            }
        ],
        bottom_callout="사람이 30분 걸리던 서식 정리를, AI가 단 40초 만에 완수합니다."
    )

    # =========================================================
    # Section 4. 내일부터 당장 써먹는 1가지 작은 실천
    # =========================================================
    # Slide 9: 실천 습관 (디플로맷 3카드: Teal, Navy, Orange)
    builder.add_pastel_3cards_slide(
        title=agenda_titles[3],
        key_message="질문 맨 끝에 '내가 원하는 완료 기준(채점표)' 딱 한 줄만 붙여보세요.",
        cards_data=[
            {
                'bg': COLOR_TEAL_CARD_BG, 'border': COLOR_TEAL_CARD_BORDER, 'title_color': COLOR_TEAL_CARD_TITLE,
                'title': '단 한 줄의 마법',
                'body': '“맨 위에 3줄 요약 넣고\n핵심은 표로 정리해줘”',
                'sub': '원하는 완료 기준(채점표) 명시'
            },
            {
                'bg': COLOR_NAVY_CARD_BG, 'border': COLOR_NAVY_CARD_BORDER, 'title_color': COLOR_NAVY_CARD_TITLE,
                'title': '자체 검증 유도',
                'body': '답변 전 자체 점검 수행\n\n기준 충족 여부 스스로 확인',
                'sub': '답변 품질의 비약적 향상'
            },
            {
                'bg': COLOR_ORANGE_CARD_BG, 'border': COLOR_ORANGE_CARD_BORDER, 'title_color': COLOR_ORANGE_CARD_TITLE,
                'title': '재작업 80% 단축',
                'body': '한 번에 결재 초안 완성\n\n반복 질문 없는 업무 완결',
                'sub': '퇴근 시간 1시간 단축'
            }
        ],
        bottom_callout="완료 기준 단 한 줄 추가가 매일 여러분의 퇴근 시간을 앞당깁니다."
    )

    # =========================================================
    # Section 5. 지속 가능한 AI 협업 워크플로우
    # =========================================================
    # Slide 10: 6단계 파이프라인 플로우 (디플로맷 테마: Teal, Navy, Ice Blue)
    builder.add_pipeline_flow_slide(
        title=agenda_titles[4],
        key_message="업무의 발견부터 자산화까지 이어지는 전사 업무 선순환 사이클",
        steps_data=[
            {
                'step': '1. 문제 정의',
                'badge_bg': COLOR_TEAL_PRIMARY,
                'text_color': COLOR_WHITE,
                'desc': '귀찮음·관성 포착\n5 Whys 본질 규명'
            },
            {
                'step': '2. 공정 분해',
                'badge_bg': COLOR_TEAL_PRIMARY,
                'text_color': COLOR_WHITE,
                'desc': '큰 단위 업무를\n작은 원자 작업 분할'
            },
            {
                'step': '3. 완료 기준',
                'badge_bg': COLOR_TEAL_PRIMARY,
                'text_color': COLOR_WHITE,
                'desc': '통과 조건 수립\n3줄 요약·양식 규격화'
            },
            {
                'step': '4. AI 실행',
                'badge_bg': COLOR_NAVY_ACCENT,
                'text_color': COLOR_WHITE,
                'desc': '맥락 재료 투입\n스킬 기반 초안 도출'
            },
            {
                'step': '5. 검토·확인',
                'badge_bg': COLOR_ICE_BLUE,
                'border_color': COLOR_TEAL_PRIMARY,
                'text_color': COLOR_TEAL_PRIMARY,
                'desc': '디렉터 팩트체크\n수치·논리 최종 검증'
            },
            {
                'step': '6. 지식 자산화',
                'badge_bg': COLOR_ICE_BLUE,
                'border_color': COLOR_TEAL_PRIMARY,
                'text_color': COLOR_TEAL_PRIMARY,
                'desc': '해결법 템플릿화\n다음 시간 50% 단축'
            }
        ],
        bottom_callout="핵심: 1회성 질문에 그치지 않고, 템플릿으로 자산화하여 복리 효율을 만듭니다."
    )

    # Slide 11: 끝맺음
    builder.add_closing_slide("경청해 주셔서 감사합니다.\n\n( 질의응답 및 자유 토론 )")

    out_file = r"D:\02_Document\20_연수회\260930-1\AI활용_한단계더_수요연수회.pptx"
    out_file_v2 = r"D:\02_Document\20_연수회\260930-1\AI활용_한단계더_수요연수회_v2.pptx"
    
    saved_paths = []
    try:
        builder.save(out_file)
        saved_paths.append(out_file)
    except PermissionError:
        print(f"[알림] '{Path(out_file).name}' 파일이 열려 있어 덮어쓰지 못했습니다.")
    
    try:
        builder.save(out_file_v2)
        saved_paths.append(out_file_v2)
    except PermissionError:
        print(f"[알림] '{Path(out_file_v2).name}' 파일이 열려 있어 덮어쓰지 못했습니다.")
        
    print(f"성공적으로 저장된 파일 목록: {saved_paths}")


if __name__ == "__main__":
    build_full_presentation()
