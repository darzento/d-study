import sys
import pptx

if sys.platform == "win32":
    sys.stdout.reconfigure(encoding='utf-8')

prs = pptx.Presentation(r"D:\02_Document\20_연수회\260930-1\AI활용_한단계더_수요연수회.pptx")
print(f"Total slides: {len(prs.slides)}")
for i, s in enumerate(prs.slides):
    title_lines = []
    title_pt = "None"
    sub_lines = []
    sub_pt = "None"
    for ph in s.placeholders:
        if ph.placeholder_format.idx == 10:
            for p in ph.text_frame.paragraphs:
                title_lines.append(p.text)
                if p.font.size:
                    title_pt = p.font.size.pt
        elif ph.placeholder_format.idx in (13, 14):
            for p in ph.text_frame.paragraphs:
                sub_lines.append(p.text)
                if p.font.size:
                    sub_pt = p.font.size.pt
    print(f"Slide {i+1:02d}:")
    print(f"  Title ({title_pt}pt): {' / '.join(title_lines)}")
    if sub_lines:
        print(f"  Sub   ({sub_pt}pt, {len(sub_lines)} lines):")
        for sl in sub_lines:
            print(f"    - {sl}")
